"""Load reference grids as strings; no numeric inference for CSV or manual entry."""
from __future__ import annotations
import csv
from datetime import date, datetime
import io
from pathlib import Path
import re


def _trim_grid(grid: list[list[str]]) -> list[list[str]]:
    width = max((len(x) for x in grid), default=0)
    if width == 0:
        return [["", ""]]
    return [row + [""] * (width - len(row)) for row in grid]


def parse_tsv(text: str) -> list[list[str]]:
    """Excel clipboard TSV, including quoted tabs/newlines and empty cells."""
    if not text:
        return []
    return _trim_grid(list(csv.reader(io.StringIO(text, newline=""), delimiter="\t")))


def read_csv(path: Path) -> list[list[str]]:
    raw = path.read_bytes()
    text = None
    for encoding in ["utf-8-sig", "utf-16", "cp950", "big5"]:
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("无法识别 CSV 编码，请另存为 UTF-8 格式")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return _trim_grid(list(csv.reader(io.StringIO(text, newline=""), dialect)))


def _xlsx_value(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        fmt = cell.number_format.strip()
        if re.fullmatch(r"0+", fmt):
            return str(value).zfill(len(fmt))
    if isinstance(value, float) and value.is_integer():
        fmt = cell.number_format.strip()
        if re.fullmatch(r"0+", fmt):
            return str(int(value)).zfill(len(fmt))
        return str(int(value))
    return str(value)


def list_sheets(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            return list(book.sheetnames)
        finally:
            book.close()
    if suffix == ".xls":
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            return book.sheet_names()
        finally:
            book.release_resources()
    return []


def read_excel(path: Path, sheet: str, header_row: int = 1) -> list[list[str]]:
    if header_row < 1:
        raise ValueError("标题行必须从第 1 行开始")
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[sheet]
            rows = [[_xlsx_value(c) for c in row] for row in ws.iter_rows(min_row=header_row)]
        finally:
            wb.close()
    elif path.suffix.lower() == ".xls":
        import xlrd
        wb = xlrd.open_workbook(str(path), on_demand=True)
        try:
            ws = wb.sheet_by_name(sheet)
            rows = [[str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)
                     for v in ws.row_values(r)]
                    for r in range(header_row - 1, ws.nrows)]
        finally:
            wb.release_resources()
    else:
        raise ValueError("只支持 .xlsx / .xls")
    return _trim_grid(rows)


def read_reference(path: Path, sheet: str = "", header_row: int = 1) -> list[list[str]]:
    if path.suffix.lower() == ".csv":
        rows = read_csv(path)
        return rows[header_row - 1:]
    return read_excel(path, sheet, header_row)


def export_grid_csv(path: Path, grid: list[list[str]]):
    # UTF-8 BOM makes Chinese headers easier to open in Excel on Windows.
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(grid)
